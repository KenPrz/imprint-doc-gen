"""Independent Office producer: python-docx 1.2.0. Synthetic, no third-party document content."""
from pathlib import Path
from io import BytesIO
import base64
from docx import Document
from docx.shared import Inches, RGBColor
out = Path(__file__).parent
host = Document()
host.add_heading('Component acceptance report', 0)
host.add_paragraph('{{>clause@1}}')
host.add_paragraph('{{>clause@1}}')
host.save(out / 'host.docx')
fragment = Document()
fragment.add_heading('Warranty clause', 1)
paragraph = fragment.add_paragraph('Service warranty ')
run = paragraph.add_run('covers defects.'); run.bold = True; run.font.color.rgb = RGBColor(0x00, 0x66, 0x99)
fragment.add_paragraph('First requirement', style='List Number')
fragment.add_paragraph('Second requirement', style='List Number')
table = fragment.add_table(rows=2, cols=2)
table.style = 'Table Grid'
table.cell(0,0).text = 'Coverage'; table.cell(0,1).text = 'Days'
table.cell(1,0).text = 'Parts'; table.cell(1,1).text = '90'
# 16x16 PNG, used to exercise real python-docx DrawingML packaging.
png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAABAAAAAQCAIAAACQkWg2AAAACXBIWXMAAAPoAAAD6AG1e1JrAAAAGUlEQVQokWNgSJtJGmIY1TBzNJTShmnSAAD42/8BAFSM9QAAAABJRU5ErkJggg==')
fragment.add_picture(BytesIO(png), width=Inches(0.15))
fragment.save(out / 'clause.docx')
